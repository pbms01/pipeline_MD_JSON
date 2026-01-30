"""
detector.py - Detecção de tipo e características do documento.

Este módulo analisa o arquivo de entrada para determinar:
- Tipo de arquivo (PDF nativo, PDF escaneado, DOCX, etc.)
- Características básicas (páginas, tamanho, etc.)
- Se contém texto extraível ou precisa de OCR/visão
"""
from pathlib import Path
from typing import Dict, Any, Optional, Tuple
import logging

try:
    import fitz  # PyMuPDF
    HAS_PYMUPDF = True
except ImportError:
    HAS_PYMUPDF = False

try:
    from docx import Document as DocxDocument
    HAS_PYTHON_DOCX = True
except ImportError:
    HAS_PYTHON_DOCX = False

from .models import DocumentType, DocumentInfo
from config.settings import (
    SUPPORTED_EXTENSIONS,
    SCANNED_PAGE_THRESHOLD,
    MAX_FILE_SIZE_MB
)

logger = logging.getLogger(__name__)


class DocumentDetector:
    """Detecta tipo e características de documentos."""

    def __init__(self):
        self.supported_extensions = SUPPORTED_EXTENSIONS

    def detect(self, file_path: Path) -> DocumentInfo:
        """
        Detecta informações do documento.

        Args:
            file_path: Caminho para o arquivo

        Returns:
            DocumentInfo com todas as informações detectadas

        Raises:
            ValueError: Se arquivo não existe ou não é suportado
        """
        file_path = Path(file_path)

        if not file_path.exists():
            raise ValueError(f"Arquivo não encontrado: {file_path}")

        extension = file_path.suffix.lower()
        if extension not in self.supported_extensions:
            raise ValueError(f"Extensão não suportada: {extension}. Suportadas: {self.supported_extensions}")

        file_size = file_path.stat().st_size
        if file_size > MAX_FILE_SIZE_MB * 1024 * 1024:
            raise ValueError(f"Arquivo muito grande: {file_size / (1024*1024):.1f}MB. Máximo: {MAX_FILE_SIZE_MB}MB")

        if extension == ".pdf":
            return self._detect_pdf(file_path, file_size)
        elif extension in (".docx", ".doc"):
            return self._detect_docx(file_path, file_size)
        else:
            return DocumentInfo(
                file_path=str(file_path),
                file_name=file_path.name,
                file_size=file_size,
                doc_type=DocumentType.UNKNOWN,
                page_count=0,
                has_text=False,
                has_images=False
            )

    def _detect_pdf(self, file_path: Path, file_size: int) -> DocumentInfo:
        """Detecta características de PDF."""
        if not HAS_PYMUPDF:
            raise ImportError("PyMuPDF (fitz) é necessário para processar PDFs. Instale com: pip install pymupdf")

        doc = fitz.open(str(file_path))

        try:
            page_count = len(doc)
            is_encrypted = doc.is_encrypted

            if is_encrypted and not doc.authenticate(""):
                return DocumentInfo(
                    file_path=str(file_path),
                    file_name=file_path.name,
                    file_size=file_size,
                    doc_type=DocumentType.PDF_NATIVE,
                    page_count=page_count,
                    has_text=False,
                    has_images=False,
                    is_encrypted=True
                )

            # Analisar algumas páginas para determinar tipo
            pages_to_check = min(5, page_count)
            text_pages = 0
            image_pages = 0
            total_text_ratio = 0.0

            for i in range(pages_to_check):
                page = doc[i]
                text = page.get_text()
                images = page.get_images()

                # Calcular ratio de texto
                page_area = page.rect.width * page.rect.height
                text_blocks = page.get_text("dict")["blocks"]
                text_area = sum(
                    (b["bbox"][2] - b["bbox"][0]) * (b["bbox"][3] - b["bbox"][1])
                    for b in text_blocks if b["type"] == 0
                )
                text_ratio = text_area / page_area if page_area > 0 else 0

                total_text_ratio += text_ratio

                if len(text.strip()) > 50:  # Tem texto significativo
                    text_pages += 1
                if images:
                    image_pages += 1

            avg_text_ratio = total_text_ratio / pages_to_check if pages_to_check > 0 else 0

            # Determinar tipo de PDF
            has_text = text_pages > 0
            has_images = image_pages > 0

            if text_pages == pages_to_check:
                doc_type = DocumentType.PDF_NATIVE
            elif text_pages == 0 and image_pages > 0:
                doc_type = DocumentType.PDF_SCANNED
            elif text_pages > 0 and image_pages > 0 and avg_text_ratio < SCANNED_PAGE_THRESHOLD:
                doc_type = DocumentType.PDF_MIXED
            else:
                doc_type = DocumentType.PDF_NATIVE

            # Extrair metadados
            metadata = {}
            if doc.metadata:
                metadata = {
                    "title": doc.metadata.get("title"),
                    "author": doc.metadata.get("author"),
                    "subject": doc.metadata.get("subject"),
                    "creator": doc.metadata.get("creator"),
                    "producer": doc.metadata.get("producer"),
                    "creation_date": doc.metadata.get("creationDate"),
                    "mod_date": doc.metadata.get("modDate"),
                }
                # Remover valores None
                metadata = {k: v for k, v in metadata.items() if v}

            return DocumentInfo(
                file_path=str(file_path),
                file_name=file_path.name,
                file_size=file_size,
                doc_type=doc_type,
                page_count=page_count,
                has_text=has_text,
                has_images=has_images,
                is_encrypted=is_encrypted,
                metadata=metadata
            )

        finally:
            doc.close()

    def _detect_docx(self, file_path: Path, file_size: int) -> DocumentInfo:
        """Detecta características de DOCX."""
        if not HAS_PYTHON_DOCX:
            raise ImportError("python-docx é necessário para processar DOCX. Instale com: pip install python-docx")

        doc = DocxDocument(str(file_path))

        # Contar parágrafos e imagens
        paragraphs = doc.paragraphs
        has_text = any(p.text.strip() for p in paragraphs)

        # Verificar imagens
        has_images = False
        for rel in doc.part.rels.values():
            if "image" in rel.reltype:
                has_images = True
                break

        # Estimar número de páginas (aproximado)
        # DOCX não tem conceito de página até ser renderizado
        char_count = sum(len(p.text) for p in paragraphs)
        estimated_pages = max(1, char_count // 2000)  # ~2000 chars por página

        # Metadados
        metadata = {}
        core_props = doc.core_properties
        if core_props:
            metadata = {
                "title": core_props.title,
                "author": core_props.author,
                "subject": core_props.subject,
                "created": str(core_props.created) if core_props.created else None,
                "modified": str(core_props.modified) if core_props.modified else None,
            }
            metadata = {k: v for k, v in metadata.items() if v}

        return DocumentInfo(
            file_path=str(file_path),
            file_name=file_path.name,
            file_size=file_size,
            doc_type=DocumentType.DOCX,
            page_count=estimated_pages,
            has_text=has_text,
            has_images=has_images,
            metadata=metadata
        )


def detect_document(file_path: Path) -> DocumentInfo:
    """
    Função de conveniência para detectar documento.

    Args:
        file_path: Caminho do arquivo

    Returns:
        DocumentInfo
    """
    detector = DocumentDetector()
    return detector.detect(file_path)


def get_document_preview(file_path: Path, max_chars: int = 5000) -> str:
    """
    Extrai preview do documento para análise rápida.

    Args:
        file_path: Caminho do arquivo
        max_chars: Máximo de caracteres a extrair

    Returns:
        Texto de preview
    """
    file_path = Path(file_path)
    extension = file_path.suffix.lower()

    if extension == ".pdf" and HAS_PYMUPDF:
        doc = fitz.open(str(file_path))
        try:
            text_parts = []
            chars_collected = 0
            for page in doc:
                if chars_collected >= max_chars:
                    break
                text = page.get_text()
                text_parts.append(text)
                chars_collected += len(text)
            return "\n\n".join(text_parts)[:max_chars]
        finally:
            doc.close()

    elif extension in (".docx", ".doc") and HAS_PYTHON_DOCX:
        doc = DocxDocument(str(file_path))
        text_parts = []
        chars_collected = 0
        for para in doc.paragraphs:
            if chars_collected >= max_chars:
                break
            text_parts.append(para.text)
            chars_collected += len(para.text)
        return "\n\n".join(text_parts)[:max_chars]

    return ""
